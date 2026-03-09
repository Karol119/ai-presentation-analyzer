import time
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

# ⏱️ Iniciar medición de tiempo
inicio = time.time()

# Ruta de tu archivo
ruta = r'C:/Users/kgonz/Desktop/TT/Presentaciones/Curso_Redes_Computadoras.pptx'

try:
    prs = Presentation(ruta)

    for i, slide in enumerate(prs.slides):
        print(f"\n=== DIAPOSITIVA {i+1} ===")
        
        # 1. Buscar el título (si existe)
        if slide.shapes.title:
            print(f"TITULO: {slide.shapes.title.text.upper()}")
        else:
            print("TITULO: (Sin título definido)")

        # 2. Buscar el resto del contenido
        print("CONTENIDO:")
        for shape in slide.shapes:
            # Evitamos repetir el título que ya imprimimos arriba
            if shape == slide.shapes.title:
                continue
                
            if hasattr(shape, "text") and shape.text.strip():
                texto_limpio = shape.text.replace('\n', ' ').strip()
                print(f" - {texto_limpio}")

except Exception as e:
    print(f"Error: {e}")

# ⏱️ Finalizar medición
fin = time.time()
print(f"\nTiempo de ejecución: {fin - inicio:.4f} segundos")