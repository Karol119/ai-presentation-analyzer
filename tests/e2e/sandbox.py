# tests/e2e/sandbox.py
"""
Sandbox aislado para la suite E2E.

Problema que resuelve
---------------------
El código productivo calcula sus rutas de datos a partir de la ubicación del
paquete `app`:

  * Base de datos:  <raiz>/storage/db/ai_analyzer.db
                    (app/data/database_manager.obtener_ruta_db)
  * Presentaciones: <raiz>/storage/presentaciones/<materia>/
                    (app/data/persistence.registrar_presentacion)

Por lo tanto, si el código se importa desde la copia temporal del sandbox, TODAS
las escrituras caen dentro del sandbox y la base de datos y los archivos del
usuario quedan intactos.

Estrategia
----------
1. Se copia `app/` a un directorio temporal.
2. Se crea `storage/db` y `storage/presentaciones` vacíos.
3. Se antepone el sandbox a `sys.path` y se purgan los módulos `app.*` ya
   importados, de modo que `import app...` resuelva a la copia.
4. `reset_db()` recrea el esquema real y siembra un catálogo mínimo.

No se modifica ni un archivo del repositorio original.
"""

from __future__ import annotations

import atexit
import os
import shutil
import sqlite3
import sys
import tempfile
from pathlib import Path

RAIZ_REPO = Path(__file__).resolve().parents[2]

_sandbox_root: Path | None = None

# ---------------------------------------------------------------------------
# Esquema real de datos (obtenido de storage/DB/ai_analyzer.db, sin datos de
# usuario). Se replica tal cual para que las FK y las columnas coincidan con
# las que usa el código productivo.
# ---------------------------------------------------------------------------
ESQUEMA = [
    """CREATE TABLE Unidad_de_Aprendizaje (
        id_unidad_aprendizaje TEXT PRIMARY KEY,
        unidad_aprendizaje TEXT NOT NULL UNIQUE,
        activa BOOLEAN NOT NULL DEFAULT 1
    )""",
    """CREATE TABLE Unidad (
        id_unidad_tematica TEXT PRIMARY KEY,
        id_unidad_aprendizaje TEXT NOT NULL,
        nombre_unidad_tematica TEXT NOT NULL,
        numero_unidad INTEGER NOT NULL,
        FOREIGN KEY (id_unidad_aprendizaje) REFERENCES Unidad_de_Aprendizaje(id_unidad_aprendizaje)
    )""",
    """CREATE TABLE Temas (
        id_tema TEXT PRIMARY KEY,
        id_unidad_tematica TEXT NOT NULL,
        nombre_tema TEXT NOT NULL,
        numero_tema INTEGER NOT NULL,
        FOREIGN KEY (id_unidad_tematica) REFERENCES Unidad(id_unidad_tematica)
    )""",
    """CREATE TABLE Subtema (
        id_subtema TEXT PRIMARY KEY,
        id_tema TEXT NOT NULL,
        nombre_subtema TEXT NOT NULL,
        numero_subtema INTEGER NOT NULL,
        embedding BLOB,
        FOREIGN KEY (id_tema) REFERENCES Temas(id_tema)
    )""",
    """CREATE TABLE Presentacion (
        id_presentacion TEXT PRIMARY KEY,
        id_unidad_aprendizaje TEXT NOT NULL,
        presentacion TEXT NOT NULL,
        visible INTEGER NOT NULL DEFAULT 1,
        FOREIGN KEY (id_unidad_aprendizaje) REFERENCES Unidad_de_Aprendizaje(id_unidad_aprendizaje)
    )""",
    """CREATE TABLE Historial_de_Versiones (
        id_version TEXT PRIMARY KEY,
        id_presentacion TEXT,
        numero_version INTEGER,
        analisis INTEGER NOT NULL DEFAULT 0,
        fecha_carga TEXT,
        reporte_tiempo TEXT,
        total_diapositivas INTEGER,
        hash TEXT,
        calificacion_presentacion REAL,
        temas_unidad_aprendizaje INTEGER,
        recomendacion_presentacion TEXT,
        ruta TEXT,
        ruta_miniatura TEXT,
        ruta_pdf TEXT,
        FOREIGN KEY (id_presentacion) REFERENCES Presentacion(id_presentacion) ON DELETE CASCADE
    )""",
    """CREATE TABLE Analisis (
        id_version TEXT NOT NULL,
        numero_diapositiva INTEGER NOT NULL,
        resultado TEXT NOT NULL,
        PRIMARY KEY (id_version, numero_diapositiva),
        FOREIGN KEY (id_version) REFERENCES Historial_de_Versiones(id_version)
    )""",
    """CREATE TABLE Presentacion_Subtema (
        id_presentacion TEXT NOT NULL,
        id_subtema TEXT NOT NULL,
        PRIMARY KEY (id_presentacion, id_subtema),
        FOREIGN KEY (id_presentacion) REFERENCES Presentacion(id_presentacion),
        FOREIGN KEY (id_subtema) REFERENCES Subtema(id_subtema)
    )""",
]

# Catálogo mínimo (equivalente a una unidad de aprendizaje del seed_temarios).
SEMILLA = {
    "ua": ("ua-prueba-0001", "Materia de Prueba", 1),
    "unidad": ("u-0001", "ua-prueba-0001", "Conceptos Fundamentales", 1),
    "tema": ("t-0001", "u-0001", "Introducción", 1),
    "subtema": ("s-0001", "t-0001", "Definiciones básicas", 1),
}

ID_MATERIA = SEMILLA["ua"][0]
NOMBRE_MATERIA = SEMILLA["ua"][1]


def bootstrap() -> Path:
    """Crea (una sola vez) el sandbox y redirige `app` hacia él. Devuelve la raíz."""
    global _sandbox_root
    if _sandbox_root is not None and _sandbox_root.exists():
        return _sandbox_root

    base_tmp = Path(os.environ.get("TEMP", tempfile.gettempdir())) / "opencode"
    base_tmp.mkdir(parents=True, exist_ok=True)
    raiz = Path(tempfile.mkdtemp(prefix="e2e-aitest-", dir=str(base_tmp)))

    # 1. Copia de la capa de aplicación (sin cachés).
    shutil.copytree(
        RAIZ_REPO / "app",
        raiz / "app",
        ignore=shutil.ignore_patterns("__pycache__", "*.pyc", "*.pyo"),
    )
    # 2. Almacenamiento propio y vacío.
    (raiz / "storage" / "db").mkdir(parents=True, exist_ok=True)
    (raiz / "storage" / "presentaciones").mkdir(parents=True, exist_ok=True)
    (raiz / "storage" / "temarios").mkdir(parents=True, exist_ok=True)

    # 3. `app` debe resolverse desde el sandbox, no desde el repositorio.
    raiz_str = str(raiz)
    while raiz_str in sys.path:
        sys.path.remove(raiz_str)
    sys.path.insert(0, raiz_str)
    for nombre in [m for m in sys.modules if m == "app" or m.startswith("app.")]:
        del sys.modules[nombre]

    _sandbox_root = raiz
    atexit.register(limpiar)
    return raiz


def limpiar() -> None:
    """Elimina el sandbox al terminar el proceso."""
    global _sandbox_root
    if _sandbox_root is not None and _sandbox_root.exists():
        shutil.rmtree(_sandbox_root, ignore_errors=True)
    _sandbox_root = None


def raiz() -> Path:
    if _sandbox_root is None:
        raise RuntimeError("sandbox.bootstrap() no ha sido invocado.")
    return _sandbox_root


def ruta_db() -> Path:
    return raiz() / "storage" / "db" / "ai_analyzer.db"


def carpeta_presentaciones() -> Path:
    return raiz() / "storage" / "presentaciones"


def conectar() -> sqlite3.Connection:
    conn = sqlite3.connect(str(ruta_db()))
    conn.execute("PRAGMA foreign_keys = ON;")
    return conn


def _vaciar_esquema() -> None:
    """Elimina todas las tablas de usuario (alternativa a borrar el archivo)."""
    conn = conectar()
    try:
        conn.execute("PRAGMA foreign_keys = OFF;")
        tablas = [f[0] for f in conn.execute(
            "SELECT name FROM sqlite_master WHERE type='table' "
            "AND name NOT LIKE 'sqlite_%'"
        ).fetchall()]
        for t in tablas:
            conn.execute(f'DROP TABLE IF EXISTS "{t}"')
        conn.commit()
    finally:
        conn.close()


def reset_db() -> None:
    """
    Recrea la base de datos con el esquema real y el catálogo mínimo.

    En Windows, si alguna conexión quedó abierta (por ejemplo tras simular un
    cierre abrupto) el archivo .db no puede borrarse; en ese caso se vacía por
    completo el esquema mediante DROP TABLE, que no requiere borrar el archivo.
    """
    import gc

    db = ruta_db()
    db.parent.mkdir(parents=True, exist_ok=True)
    if db.exists():
        for _intento in range(5):
            try:
                db.unlink()
                break
            except PermissionError:
                gc.collect()
        else:
            _vaciar_esquema()

    conn = conectar()
    try:
        for ddl in ESQUEMA:
            conn.execute(ddl)
        ua = SEMILLA["ua"]
        un = SEMILLA["unidad"]
        te = SEMILLA["tema"]
        su = SEMILLA["subtema"]
        conn.execute("INSERT INTO Unidad_de_Aprendizaje VALUES (?, ?, ?)", ua)
        conn.execute("INSERT INTO Unidad VALUES (?, ?, ?, ?)", un)
        conn.execute("INSERT INTO Temas VALUES (?, ?, ?, ?)", te)
        conn.execute("INSERT INTO Subtema (id_subtema, id_tema, nombre_subtema, numero_subtema) VALUES (?, ?, ?, ?)", su)
        conn.commit()
    finally:
        conn.close()
    limpiar_presentaciones()


def limpiar_presentaciones() -> None:
    carpeta = carpeta_presentaciones()
    if carpeta.exists():
        for hijo in carpeta.iterdir():
            if hijo.is_dir():
                shutil.rmtree(hijo, ignore_errors=True)
            else:
                try:
                    hijo.unlink()
                except OSError:
                    pass


# ---------------------------------------------------------------------------
# Datos de prueba: presentaciones PPTX sintéticas
# ---------------------------------------------------------------------------

def crear_pptx(destino: Path, n_slides: int = 4, titulo: str = "Presentación de prueba") -> Path:
    """Genera un .pptx válido y minimo con `n_slides` diapositivas de contenido."""
    from pptx import Presentation

    prs = Presentation()
    layout = prs.slide_layouts[1]  # Título y contenido
    for i in range(n_slides):
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = f"{titulo} — Tema {i + 1}"
        cuerpo = slide.placeholders[1].text_frame
        cuerpo.text = (
            "Una red de computadoras conecta dispositivos mediante protocolos "
            "estandarizados que permiten el intercambio de informacion."
        )
        cuerpo.add_paragraph().text = (
            "El modelo de referencia describe niveles de abstraccion para "
            "separar las funciones de transmision y recepcion de datos."
        )
    destino.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(destino))
    return destino


def crear_archivo_pesado(destino: Path, mb: int = 31) -> Path:
    """Crea un archivo de tamaño controlado para probar el limite de 30 MB."""
    destino.parent.mkdir(parents=True, exist_ok=True)
    with open(destino, "wb") as f:
        f.seek(mb * 1024 * 1024 - 1)
        f.write(b"\0")
    return destino


def crear_pdf_minimo(destino: Path, paginas: int = 4) -> Path:
    """Crea un PDF minimo (para el visor de diapositivas) sin depender de PowerPoint."""
    import pymupdf

    destino.parent.mkdir(parents=True, exist_ok=True)
    doc = pymupdf.open()
    for i in range(paginas):
        pagina = doc.new_page()
        pagina.insert_text((72, 72), f"Diapositiva {i + 1}")
    doc.save(str(destino))
    doc.close()
    return destino
