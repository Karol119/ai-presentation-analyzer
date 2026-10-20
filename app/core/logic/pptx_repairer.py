# app/core/logic/pptx_repairer.py
"""
Reparación de presentaciones PPTX dañadas (filosofía PowerPoint).

PowerPoint, al abrir un archivo con problemas, lo "repara" quitando las
partes ilegibles y conservando todo lo legible. Este módulo replica esa
misma política de forma programática, en DOS fases:

  FASE 1 (a nivel paquete/zip):
    Quita del paquete las diapositivas cuyo XML los motores de lectura no
    pueden parsear (típico cuando el archivo no se guardó/sincronizó bien
    y PowerPoint pide "reparar"). python-pptx no puede abrir un paquete
    con XML ilegible, así que esto se hace con zip + lxml:
      - elimina ppt/slides/slideN.xml dañados
      - elimina sus relaciones de ppt/_rels/presentation.xml.rels
      - elimina sus p:sldId de ppt/presentation.xml
    Esto es EXACTAMENTE lo que PowerPoint hace: conserva el resto.

  FASE 2 (con python-pptx):
    Recorre cada diapositiva de la misma forma que el extractor del
    análisis (título, shapes, párrafos, grupos, imágenes incrustadas):
      - shape ilegible -> se elimina ese elemento
      - diapositiva inservible -> se elimina completa
    y guarda un NUEVO archivo (el original queda intacto).

La salida la vuelve a validar el orquestador antes de registrarla, de modo
que el usuario ya nunca se encontrará con el crasheo al analizar.
"""
import os
import zipfile
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

_NS_P        = "http://schemas.openxmlformats.org/presentationml/2006/main"
_NS_R        = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_NS_REL      = "http://schemas.openxmlformats.org/package/2006/relationships"
_RT_SLIDE    = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide"
_ROOT_SLIDES = "ppt/slides/"


# ---------------------------------------------------------------------------
# FASE 1 · Limpieza a nivel paquete (diapositivas con XML ilegible)
# ---------------------------------------------------------------------------

def _xml_tiene_sintaxis_ilegible(data: bytes) -> bool:
    try:
        etree.fromstring(data)
        return False
    except Exception:
        return True


def _fase_1_quitar_diapositivas_xml_ilegibles(ruta_origen: str, ruta_salida: str) -> int:
    """
    Reconstruye el paquete excluyendo las diapositivas cuyo XML no se puede
    parsear, actualizando presentation.xml y sus relaciones. Devuelve cuántas
    diapositivas se quitaron (0 = nada que reparar a este nivel).
    """
    with zipfile.ZipFile(ruta_origen) as paquete:
        nombres = paquete.namelist()

        # 1. Detectar diapositivas con XML ilegible
        diapositivas_malas = {
            nombre for nombre in nombres
            if nombre.startswith(_ROOT_SLIDES)
            and nombre.endswith(".xml")
            and "/_rels/" not in nombre
            and _xml_tiene_sintaxis_ilegible(paquete.read(nombre))
        }
        if not diapositivas_malas:
            return 0

        bases_malas = {ruta.rsplit("/", 1)[-1] for ruta in diapositivas_malas}

        # 2. Quitar de presentation.xml.rels las relaciones hacia esas slides
        nombre_rels = "ppt/_rels/presentation.xml.rels"
        ids_quitados = set()
        doc_rels = None
        if nombre_rels in nombres:
            doc_rels = etree.fromstring(paquete.read(nombre_rels))
            for rel in list(doc_rels):
                tipo = rel.get("Type", "")
                objetivo = rel.get("Target", "")
                base = objetivo.rsplit("/", 1)[-1]
                if tipo == _RT_SLIDE and base in bases_malas:
                    ids_quitados.add(rel.get("Id"))
                    doc_rels.remove(rel)

        # 3. Quitar de presentation.xml los sldId referidos a esas relaciones
        doc_pres = etree.fromstring(paquete.read("ppt/presentation.xml"))
        lista_sld = doc_pres.find(f"{{{_NS_P}}}sldIdLst")
        if lista_sld is not None:
            for sld in list(lista_sld):
                if sld.get(f"{{{_NS_R}}}id") in ids_quitados:
                    lista_sld.remove(sld)

        # 4. Reescribir el paquete
        with zipfile.ZipFile(ruta_salida, "w", zipfile.ZIP_DEFLATED) as salida:
            for info in paquete.infolist():
                nombre = info.filename
                if nombre in diapositivas_malas:
                    continue  # parte dañada eliminada
                # Rels huérfanas de las diapositivas eliminadas
                if nombre.startswith("ppt/slides/_rels/"):
                    nombre_rels_slide = nombre.rsplit("/", 1)[-1]
                    slide_referida = "ppt/slides/" + nombre_rels_slide[: -len(".rels")]
                    if slide_referida in diapositivas_malas:
                        continue
                if nombre == "ppt/presentation.xml":
                    salida.writestr(
                        info,
                        etree.tostring(doc_pres, xml_declaration=True,
                                       encoding="UTF-8", standalone=True),
                    )
                elif nombre == nombre_rels and doc_rels is not None:
                    salida.writestr(
                        info,
                        etree.tostring(doc_rels, xml_declaration=True,
                                       encoding="UTF-8", standalone=True),
                    )
                else:
                    salida.writestr(info, paquete.read(nombre))

    return len(diapositivas_malas)


# ---------------------------------------------------------------------------
# FASE 2 · Limpieza con python-pptx (los mismos accesos del extractor)
# ---------------------------------------------------------------------------

def _es_imagen_vinculada(exc) -> bool:
    """True si la excepción corresponde a una imagen vinculada/externa,
    que PowerPoint abre bien y no es corrupción."""
    return "no embedded image" in str(exc).lower()


def _leer_shape(shape):
    """
    Lee un shape exactamente como lo hará el extractor de texto
    (shape_type, texto de párrafos, grupos anidados, imagen incrustada).
    Lanza excepción si el contenido está dañado.
    """
    _ = shape.shape_type
    if shape.has_text_frame:
        for parrafo in shape.text_frame.paragraphs:
            _ = parrafo.text
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for hijo in list(shape.shapes):
            _leer_shape(hijo)
    if getattr(type(shape), "image", None) is not None:
        try:
            _ = shape.image.blob  # fuerza la lectura de la imagen incrustada
        except ValueError as exc:
            if not _es_imagen_vinculada(exc):
                raise


def _quitar_shape(shape):
    """Elimina el shape del árbol XML de su diapositiva/partes base."""
    elemento = shape._element
    padre = elemento.getparent()
    if padre is not None:
        padre.remove(elemento)
    # Si era una imagen incrustada, retirar también su relación para que la
    # imagen dañada no se guarde dentro del archivo reparado.
    try:
        rId = elemento.blip_rId
        if rId:
            shape.part.drop_rel(rId)
    except Exception:
        pass


def _quitar_diapositiva(presentacion, sld_id):
    """Elimina una diapositiva del paquete (relación + lista de slides)."""
    try:
        presentacion.part.drop_rel(sld_id.rId)
    except Exception:
        pass
    presentacion.slides._sldIdLst.remove(sld_id)


def _limpiar_diapositiva(diapo) -> int:
    """Devuelve cuántos shapes ilegibles se quitaron de la diapositiva."""
    quitados = 0
    for shape in list(diapo.shapes):
        try:
            _leer_shape(shape)
        except Exception:
            _quitar_shape(shape)
            quitados += 1
    return quitados


def _fase_2_limpiar_con_python_pptx(ruta_origen: str, ruta_salida: str, resultado: dict):
    """Abre el paquete y quita shapes/diapositivas ilegibles guardando una
    copia limpia en ruta_salida."""
    presentacion = Presentation(ruta_origen)

    # Pareja (entrada de sldIdLst, diapositiva) en el mismo orden: sirve para
    # localizar y quitar diapositivas enteras si se necesita.
    sld_ids = list(presentacion.slides._sldIdLst)
    diapositivas = list(presentacion.slides)
    mapa = dict(zip(sld_ids, diapositivas))

    for sld_id, diapo in mapa.items():
        try:
            _ = diapo.shapes.title          # el análisis lo lee primero
            quitados_shape = _limpiar_diapositiva(diapo)
            resultado["elementos_eliminados"] += quitados_shape
            resultado["diapositivas_conservadas"] += 1
        except Exception:
            # La diapositiva entera es ilegible -> se elimina completa
            _quitar_diapositiva(presentacion, sld_id)
            resultado["diapositivas_eliminadas"] += 1

    presentacion.save(ruta_salida)


# ---------------------------------------------------------------------------
# API pública
# ---------------------------------------------------------------------------

def reparar_pptx(ruta_origen: str, ruta_salida: str) -> dict:
    """
    Repara el archivo conservando todo lo legible y quitando lo dañado.
    Devuelve un resumen:
        {
          "exito": bool,
          "diapositivas_eliminadas": int,
          "elementos_eliminados": int,
          "diapositivas_conservadas": int,
          "mensaje": str  (vacío si exito)
        }
    """
    resultado = {
        "exito": False,
        "diapositivas_eliminadas": 0,
        "elementos_eliminados": 0,
        "diapositivas_conservadas": 0,
        "mensaje": "",
    }

    ruta_paso = None
    try:
        # FASE 1 · quitar diapositivas con XML ilegible (si aplica)
        quitadas_fase1 = 0
        try:
            ruta_paso = ruta_salida + ".paso1"
            quitadas_fase1 = _fase_1_quitar_diapositivas_xml_ilegibles(
                ruta_origen, ruta_paso
            )
            if quitadas_fase1:
                resultado["diapositivas_eliminadas"] += quitadas_fase1
        except Exception as e:
            print(f"[pptx_repairer] Fase 1 (zip) sin efecto: {e}")
            quitadas_fase1 = 0

        # FASE 2 · limpieza de shapes/diapositivas ilegibles
        origen = ruta_paso if quitadas_fase1 else ruta_origen
        _fase_2_limpiar_con_python_pptx(origen, ruta_salida, resultado)
        resultado["exito"] = True

    except Exception as e:
        resultado["mensaje"] = str(e)
    finally:
        if ruta_paso and os.path.exists(ruta_paso):
            try:
                os.remove(ruta_paso)
            except OSError:
                pass
    return resultado
