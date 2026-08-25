# app/core/logic/pptx_modifier.py
import os
import zipfile
import tempfile
import shutil
import xml.etree.ElementTree as ET
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

def aplicar_mejoras_pptx(ruta_original: str, ruta_destino: str, datos_analisis: dict) -> tuple[bool, str]:
    """
    Punto de entrada de la capa lógica. Orquesta la aplicación de mejoras en lote.
    Copia el archivo original y aplica las divisiones/reemplazos de texto
    en orden inverso para no afectar los índices XML.
    """
    try:
        # 1. Asegurar directorio y crear una copia exacta de trabajo
        os.makedirs(os.path.dirname(ruta_destino), exist_ok=True)
        shutil.copyfile(ruta_original, ruta_destino)

        diapositivas_datos = datos_analisis.get("slides", [])

        # 2. Recorrer de atrás hacia adelante (¡Clave para no arruinar los índices!)
        for info_diapositiva in reversed(diapositivas_datos):
            if info_diapositiva.get("requiere_reestructuracion") and info_diapositiva.get("reestructuracion"):
                
                num_diapositiva = info_diapositiva["slide_number"] # Base 1
                sugerencias = info_diapositiva["reestructuracion"].get("diapositivas_generadas", [])

                # Empaquetar los textos generados por la IA para el procesador
                lista_contenidos = []
                for sug in sugerencias:
                    lista_contenidos.append({
                        "titulo": sug.get("titulo_sugerido", ""),
                        "cuerpo": sug.get("contenido_optimizado", "")
                    })

                # 3. Aplicar la clonación y reemplazo físico en el archivo destino
                if lista_contenidos:
                    _dividir_y_reemplazar_diapositiva(ruta_destino, num_diapositiva, lista_contenidos)

        return True, "Presentación optimizada generada con éxito."
    except Exception as e:
        print(f"Error al manipular el PPTX (OpenXML): {e}")
        return False, f"Error al procesar la presentación: {str(e)}"


# =============================================================================
# FUNCIONES DE MANIPULACIÓN OPENXML Y PPTX (NÚCLEO DEL CLONADOR)
# =============================================================================

def _obtener_total_diapositivas(ruta_pptx: str) -> int:
    """Devuelve el total de diapositivas usando python-pptx."""
    presentacion = Presentation(ruta_pptx)
    return len(presentacion.slides)

def _obtener_etiqueta_local(etiqueta: str) -> str:
    """Limpia los namespaces de las etiquetas XML para facilitar la búsqueda."""
    if '}' in etiqueta:
        return etiqueta.split('}', 1)[1]
    return etiqueta

def _duplicar_diapositiva_en_paquete(ruta_pptx: str, indice_objetivo: int, num_copias: int):
    """
    Duplicación limpia y nativa a nivel de paquete OpenXML físico (ZIP).
    Clona los archivos .xml internos para preservar fondos y másteres.
    """
    if num_copias <= 1:
        return ruta_pptx

    directorio_temp = tempfile.mkdtemp()
    try:
        # 1. Descomprimir PPTX como si fuera un ZIP
        with zipfile.ZipFile(ruta_pptx, 'r') as zip_ref:
            zip_ref.extractall(directorio_temp)

        # Rutas a los manifiestos XML principales
        ruta_pres_xml = os.path.join(directorio_temp, "ppt", "presentation.xml")
        ruta_pres_rels = os.path.join(directorio_temp, "ppt", "_rels", "presentation.xml.rels")
        ruta_content_types = os.path.join(directorio_temp, "[Content_Types].xml")

        arbol_pres = ET.parse(ruta_pres_xml)
        raiz_pres = arbol_pres.getroot()

        arbol_rels = ET.parse(ruta_pres_rels)
        raiz_rels = arbol_rels.getroot()

        arbol_ct = ET.parse(ruta_content_types)
        raiz_ct = arbol_ct.getroot()

        # 2. Localizar la lista de diapositivas
        lista_ids_diapositivas = None
        for elemento in raiz_pres.iter():
            if _obtener_etiqueta_local(elemento.tag) == "sldIdLst":
                lista_ids_diapositivas = elemento
                break

        if lista_ids_diapositivas is None:
            raise ValueError("No se encontró la lista de diapositivas en presentation.xml")

        elementos_id_diapositiva = [hijo for hijo in lista_ids_diapositivas if _obtener_etiqueta_local(hijo.tag) == "sldId"]
        elemento_objetivo = elementos_id_diapositiva[indice_objetivo]

        id_relacion_objetivo = None
        for clave_attr, valor_attr in elemento_objetivo.attrib.items():
            if clave_attr.endswith("id") and valor_attr.startswith("rId"):
                id_relacion_objetivo = valor_attr
                break

        # 3. Localizar las dependencias de la diapositiva objetivo
        relacion_objetivo = None
        for relacion in raiz_rels.iter():
            if _obtener_etiqueta_local(relacion.tag) == "Relationship":
                if relacion.attrib.get("Id") == id_relacion_objetivo:
                    relacion_objetivo = relacion
                    break

        archivo_diapositiva_objetivo = os.path.basename(relacion_objetivo.attrib.get("Target", ""))
        ruta_diapositiva_objetivo = os.path.join(directorio_temp, "ppt", "slides", archivo_diapositiva_objetivo)
        ruta_rels_objetivo = os.path.join(directorio_temp, "ppt", "slides", "_rels", f"{archivo_diapositiva_objetivo}.rels")

        # 4. Calcular identificadores máximos actuales para no causar conflictos
        dir_diapositivas = os.path.join(directorio_temp, "ppt", "slides")
        archivos_existentes = [f for f in os.listdir(dir_diapositivas) if f.startswith("slide") and f.endswith(".xml")]
        
        numeros_diapositivas = [int(f.replace("slide", "").replace(".xml", "")) for f in archivos_existentes if f.replace("slide", "").replace(".xml", "").isdigit()]
        max_num_diapositiva = max(numeros_diapositivas) if numeros_diapositivas else 0

        ids_relaciones = []
        for rel in raiz_rels.iter():
            if _obtener_etiqueta_local(rel.tag) == "Relationship":
                str_rid = rel.attrib.get("Id", "")
                if str_rid.startswith("rId") and str_rid[3:].isdigit():
                    ids_relaciones.append(int(str_rid[3:]))
        max_rid = max(ids_relaciones) if ids_relaciones else 0

        ids_sld = [int(elem.attrib.get("id")) for elem in elementos_id_diapositiva if elem.attrib.get("id", "").isdigit()]
        max_id_sld = max(ids_sld) if ids_sld else 255

        # Espacios de nombres (Namespaces) obligatorios de OpenXML
        ns_rels = "http://schemas.openxmlformats.org/package/2006/relationships"
        ns_p = "http://schemas.openxmlformats.org/presentationml/2006/main"
        ns_r = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
        ns_ct = "http://schemas.openxmlformats.org/package/2006/content-types"

        indice_insercion = list(lista_ids_diapositivas).index(elemento_objetivo) + 1

        # 5. Generar copias físicas en el directorio temporal
        for i in range(1, num_copias):
            max_num_diapositiva += 1
            max_rid += 1
            max_id_sld += 1

            nuevo_archivo = f"slide{max_num_diapositiva}.xml"
            nueva_ruta_xml = os.path.join(directorio_temp, "ppt", "slides", nuevo_archivo)
            nueva_ruta_rels = os.path.join(directorio_temp, "ppt", "slides", "_rels", f"{nuevo_archivo}.rels")
            nuevo_rid = f"rId{max_rid}"
            nuevo_id_str = str(max_id_sld)

            shutil.copyfile(ruta_diapositiva_objetivo, nueva_ruta_xml)
            if os.path.exists(ruta_rels_objetivo):
                shutil.copyfile(ruta_rels_objetivo, nueva_ruta_rels)

            # Registrar en content_types
            elemento_override = ET.Element(f"{{{ns_ct}}}Override", {
                "PartName": f"/ppt/slides/{nuevo_archivo}",
                "ContentType": "application/vnd.openxmlformats-officedocument.presentationml.slide+xml"
            })
            raiz_ct.append(elemento_override)

            # Registrar relación
            elemento_relacion = ET.Element(f"{{{ns_rels}}}Relationship", {
                "Id": nuevo_rid,
                "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide",
                "Target": f"slides/{nuevo_archivo}"
            })
            raiz_rels.append(elemento_relacion)

            # Insertar en la presentación visible
            nuevo_elemento_sldid = ET.Element(f"{{{ns_p}}}sldId", {
                "id": nuevo_id_str,
                f"{{{ns_r}}}id": nuevo_rid
            })
            lista_ids_diapositivas.insert(indice_insercion, nuevo_elemento_sldid)
            indice_insercion += 1

        # 6. Guardar manifiestos modificados
        arbol_pres.write(ruta_pres_xml, encoding="utf-8", xml_declaration=True)
        arbol_rels.write(ruta_pres_rels, encoding="utf-8", xml_declaration=True)
        arbol_ct.write(ruta_content_types, encoding="utf-8", xml_declaration=True)

        ruta_respaldo = ruta_pptx + ".bak"
        shutil.copyfile(ruta_pptx, ruta_respaldo)

        # 7. Volver a comprimir a .pptx
        with zipfile.ZipFile(ruta_pptx, 'w', zipfile.ZIP_DEFLATED) as zip_salida:
            for nombre_carpeta, _, nombres_archivos in os.walk(directorio_temp):
                for nombre_archivo in nombres_archivos:
                    ruta_archivo = os.path.join(nombre_carpeta, nombre_archivo)
                    nombre_relativo = os.path.relpath(ruta_archivo, directorio_temp)
                    zip_salida.write(ruta_archivo, nombre_relativo)

        if os.path.exists(ruta_respaldo):
            os.remove(ruta_respaldo)

    finally:
        shutil.rmtree(directorio_temp, ignore_errors=True)

def _obtener_formas_texto_recursivo(coleccion_formas) -> list:
    """Extrae recursivamente todas las formas con texto, evaluando grupos."""
    formas_texto = []
    for forma in coleccion_formas:
        if forma.shape_type == MSO_SHAPE_TYPE.GROUP:
            formas_texto.extend(_obtener_formas_texto_recursivo(forma.shapes))
        elif forma.has_text_frame:
            formas_texto.append(forma)
    return formas_texto

def _extraer_estilo_fuente(forma) -> dict:
    """Extrae con precisión las propiedades tipográficas existentes."""
    estilo = {"name": None, "size": None, "bold": None, "italic": None, "color_rgb": None}
    if forma.has_text_frame and forma.text_frame.paragraphs:
        for parrafo in forma.text_frame.paragraphs:
            for run in parrafo.runs:
                if run.text.strip():
                    estilo["name"] = run.font.name
                    estilo["size"] = run.font.size
                    estilo["bold"] = run.font.bold
                    estilo["italic"] = run.font.italic
                    try:
                        if run.font.color and run.font.color.rgb:
                            estilo["color_rgb"] = run.font.color.rgb
                    except Exception:
                        pass
                    return estilo
    return estilo

def _establecer_texto_preservando_estilo(forma, texto: str):
    """Sobreescribe un cuadro de texto existente manteniendo su tipografía original."""
    estilo = _extraer_estilo_fuente(forma)
    marco_texto = forma.text_frame
    
    # Mantener únicamente el primer párrafo y eliminar residuales
    while len(marco_texto.paragraphs) > 1:
        parrafo_extra = marco_texto.paragraphs[-1]._p
        parrafo_extra.getparent().remove(parrafo_extra)
        
    parrafo = marco_texto.paragraphs[0]
    parrafo.text = ""
    run = parrafo.add_run()
    run.text = texto
    
    if estilo.get("name"): run.font.name = estilo["name"]
    if estilo.get("size"): run.font.size = estilo["size"]
    if estilo.get("bold") is not None: run.font.bold = estilo["bold"]
    if estilo.get("italic") is not None: run.font.italic = estilo["italic"]
    if estilo.get("color_rgb"): run.font.color.rgb = estilo["color_rgb"]

def _actualizar_diapositiva_jerarquica(diapositiva, texto_titulo: str, texto_cuerpo: str):
    """
    Localiza cajas de texto (ordenadas verticalmente).
    Reemplaza la primera con el Título, la segunda con el Cuerpo y vacía residuos.
    """
    todas_las_formas = _obtener_formas_texto_recursivo(diapositiva.shapes)
    
    # Filtrar formas que tengan texto real y ordenarlas de arriba hacia abajo
    formas_con_contenido = [f for f in todas_las_formas if f.text_frame.text.strip()]
    formas_con_contenido.sort(key=lambda f: f.top)

    if not formas_con_contenido:
        return

    # Si la diapositiva original tenía al menos 2 cajas (Título y Cuerpo)
    if len(formas_con_contenido) >= 2:
        if texto_titulo:
            _establecer_texto_preservando_estilo(formas_con_contenido[0], texto_titulo)
        else:
            formas_con_contenido[0].text_frame.text = ""

        if texto_cuerpo:
            _establecer_texto_preservando_estilo(formas_con_contenido[1], texto_cuerpo)
        else:
            formas_con_contenido[1].text_frame.text = ""

        # Limpiar notas o subtítulos sobrantes del diseño original
        for forma_extra in formas_con_contenido[2:]:
            forma_extra.text_frame.text = ""
            
    # Si la diapositiva original tenía solo 1 gran caja de texto
    elif len(formas_con_contenido) == 1:
        texto_completo = f"{texto_titulo}\n{texto_cuerpo}".strip() if texto_titulo else texto_cuerpo
        _establecer_texto_preservando_estilo(formas_con_contenido[0], texto_completo)

def _dividir_y_reemplazar_diapositiva(ruta_pptx: str, num_diapositiva_objetivo: int, lista_contenidos: list):
    """
    Orquesta la clonación a nivel de paquete y la inyección jerárquica de texto.
    Abre y guarda el archivo en la misma ruta provista.
    """
    total = _obtener_total_diapositivas(ruta_pptx)
    if num_diapositiva_objetivo < 1 or num_diapositiva_objetivo > total:
        raise ValueError(f"Número de diapositiva inválido (1 a {total}).")

    indice_objetivo = num_diapositiva_objetivo - 1
    num_copias = len(lista_contenidos)

    # 1. Duplicación estricta a nivel XML ZIP
    _duplicar_diapositiva_en_paquete(ruta_pptx, indice_objetivo, num_copias)

    # 2. Inyección limpia de los textos sugeridos por la IA
    presentacion = Presentation(ruta_pptx)
    for i, contenido in enumerate(lista_contenidos):
        diapositiva_actual = presentacion.slides[indice_objetivo + i]
        _actualizar_diapositiva_jerarquica(
            diapositiva_actual,
            texto_titulo=contenido.get("titulo", ""),
            texto_cuerpo=contenido.get("cuerpo", "")
        )

    # Sobreescribimos el archivo (que ya es el archivo destino temporal)
    presentacion.save(ruta_pptx)